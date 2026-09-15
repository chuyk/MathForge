import os
import re
import docx
from docx import Document
from docx.shared import Pt, Inches, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.parts.image import ImagePart
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT

# 匯入 OMML 核心模組
try:
    from _tools.omml_helper import add_exam_runs, set_run_font, sanitize_text, configure_document_normal_style, ensure_document_font_consistency
except ImportError:
    from omml_helper import add_exam_runs, set_run_font, sanitize_text, configure_document_normal_style, ensure_document_font_consistency

# 規範色彩常數
COLOR_BLACK = RGBColor(0, 0, 0)
COLOR_RED = RGBColor(180, 0, 0)
COLOR_NAVY = RGBColor(20, 50, 110)
COLOR_GRAY = RGBColor(140, 140, 140)
COLOR_SUBTITLE = RGBColor(80, 80, 80)

def configure_b4_section(section):
    """設定標準 JIS B4 紙張與 22mm 邊界"""
    section.page_width = Mm(257)
    section.page_height = Mm(364)
    section.top_margin = Mm(22)
    section.bottom_margin = Mm(22)
    section.left_margin = Mm(22)
    section.right_margin = Mm(22)

def upgrade_to_modern_word_mode(doc):
    """
    將 Word 文件升級為最新現代模式（Word 2016/2019/2021/365），徹底移除「相容模式」限制。
    確保 Word 原生向量 SVG 圖形、OMML 數學公式與現代繪圖物件能以 100% 完整功能渲染，
    並原生解鎖「圖形格式」與「轉換為圖形 (Convert to Shape)」可編輯功能。
    """
    try:
        settings = doc.settings.element
        compat_nodes = settings.xpath('.//w:compatSetting[@w:name="compatibilityMode"]')
        for node in compat_nodes:
            node.set(qn('w:val'), '15')
    except Exception:
        pass

def add_runs_paragraph(doc, runs_data, line_spacing=1.2, space_before=1, space_after=3, left_indent=0, force_color=None, keep_with_next=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = keep_with_next
    if left_indent > 0:
        p.paragraph_format.left_indent = Inches(left_indent)
    add_exam_runs(p, runs_data, default_font="標楷體", ascii_font="Times New Roman", default_size=13, force_color=force_color)
    return p

def add_clean_paragraph(doc, space_before=0, space_after=2, line_spacing=1.15, keep_with_next=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.keep_with_next = keep_with_next
    return p

def add_exam_image(paragraph, png_path, svg_path=None, width_inch=3.2):
    """
    在段落中加入試題附圖。
    若提供 svg_path 且存在，則以微軟 OOXML 標準嵌入 SVG 向量圖，並以 png_path 作為向下相容 fallback。
    老師在 Word 中點選圖片後，可點選「圖形格式」->「轉換為圖形 (Convert to Shape)」拆解為原生向量形狀與文字進行編輯！
    """
    run = paragraph.add_run()
    inline_shape = run.add_picture(png_path, width=Inches(width_inch))
    
    if svg_path and os.path.exists(svg_path):
        try:
            with open(svg_path, 'rb') as f:
                svg_bytes = f.read()
            
            doc_part = paragraph.part
            package = doc_part.package
            
            # 檢查 package 中是否已存在相同 SVG part，避免重複打包膨脹
            existing_part = None
            for p in package.parts:
                if str(p.partname).endswith('.svg') and getattr(p, 'blob', None) == svg_bytes:
                    existing_part = p
                    break
            
            if existing_part is not None:
                svg_part = existing_part
            else:
                svg_count = sum(1 for p in package.parts if str(p.partname).endswith('.svg'))
                partname = PackURI(f'/word/media/diagram_{svg_count + 1}.svg')
                svg_part = ImagePart(partname, 'image/svg+xml', svg_bytes)
                package.parts.append(svg_part)
            
            rId_svg = doc_part.relate_to(svg_part, RT.IMAGE)
            
            blips = inline_shape._inline.xpath('.//a:blip')
            if blips:
                extLst_xml = f'''<a:extLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
                    <a:ext uri="{{28A0092B-C50C-407E-A947-70E740481C1C}}">
                        <a14:useLocalDpi xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" val="0"/>
                    </a:ext>
                    <a:ext uri="{{96DAC541-7B7A-43D3-8B79-37D633B846F1}}">
                        <asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:embed="{rId_svg}"/>
                    </a:ext>
                </a:extLst>'''
                blips[0].append(parse_xml(extLst_xml))
        except Exception as e:
            # 若 SVG 附加發生異常，維持標準 PNG 正常顯示，不影響試卷產出
            pass
    return inline_shape

def render_options(doc, options, options_type="text"):
    """
    三層式智慧選項排版：
    1. 單行四欄 (text, 1 row)
    2. 雙行雙欄 (2x2, Tab Stop 4.2 英吋，完全無表格)
    3. 單選單行四行直列 (1x4, lines, 長選項 > 20 字)
    """
    if options_type in ("1x4", "lines", "column"):
        for opt_text in options:
            add_runs_paragraph(doc, [opt_text], space_before=1, space_after=3, left_indent=0.35)
    elif options_type == "2x2":
        # 檢測是否有選項字數過長 (> 20 字且無公式)
        flat_opts = []
        if isinstance(options, list):
            for item in options:
                if isinstance(item, (list, tuple)):
                    flat_opts.extend(item)
                else:
                    flat_opts.append(item)
        
        is_too_long = any(len(c) > 20 for c in flat_opts if not ('$' in c and len(c) < 35))
        if is_too_long or len(flat_opts) != 4:
            for opt_text in flat_opts:
                add_runs_paragraph(doc, [opt_text], space_before=1, space_after=3, left_indent=0.35)
        else:
            # 第一行 (A) 與 (B)
            p1 = add_clean_paragraph(doc, space_before=1, space_after=1, keep_with_next=True)
            p1.paragraph_format.left_indent = Inches(0.35)
            p1.paragraph_format.tab_stops.add_tab_stop(Inches(4.2), WD_TAB_ALIGNMENT.LEFT)
            add_exam_runs(p1, [flat_opts[0]], default_font="標楷體", ascii_font="Times New Roman", default_size=13)
            p1.add_run("\t")
            add_exam_runs(p1, [flat_opts[1]], default_font="標楷體", ascii_font="Times New Roman", default_size=13)

            # 第二行 (C) 與 (D)
            p2 = add_clean_paragraph(doc, space_before=1, space_after=3)
            p2.paragraph_format.left_indent = Inches(0.35)
            p2.paragraph_format.tab_stops.add_tab_stop(Inches(4.2), WD_TAB_ALIGNMENT.LEFT)
            add_exam_runs(p2, [flat_opts[2]], default_font="標楷體", ascii_font="Times New Roman", default_size=13)
            p2.add_run("\t")
            add_exam_runs(p2, [flat_opts[3]], default_font="標楷體", ascii_font="Times New Roman", default_size=13)
    else:
        # 單行四欄 (text)
        if isinstance(options, list) and len(options) == 1:
            add_runs_paragraph(doc, [options[0]], space_before=1, space_after=3, left_indent=0.35)
        elif isinstance(options, list):
            line_str = "        ".join(options)
            add_runs_paragraph(doc, [line_str], space_before=1, space_after=3, left_indent=0.35)
        else:
            add_runs_paragraph(doc, [str(options)], space_before=1, space_after=3, left_indent=0.35)

def build_student_exam(questions_data, title, subtitle, out_path):
    """產出檔案 1：全卷試題與後附詳解(B4_13pt版).docx"""
    doc = Document()
    configure_document_normal_style(doc, size_pt=13)
    upgrade_to_modern_word_mode(doc)
    configure_b4_section(doc.sections[0])

    # 標題
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run(title)
    set_run_font(r_title, font_name="標楷體", ascii_font="Times New Roman", size_pt=18, bold=True, color_rgb=COLOR_NAVY)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run(subtitle)
    set_run_font(r_sub, font_name="標楷體", ascii_font="Times New Roman", size_pt=12, color_rgb=COLOR_SUBTITLE)

    # 全部試題
    for q in questions_data:
        add_runs_paragraph(doc, [q["stem"]], space_before=2, space_after=3, keep_with_next=True)
        
        # 題目附圖 (下一行純流式呈現)
        if q.get("image_path") and os.path.exists(q["image_path"]):
            p_img = doc.add_paragraph()
            p_img.paragraph_format.left_indent = Inches(0.4)
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.paragraph_format.keep_with_next = True
            add_exam_image(p_img, q["image_path"], q.get("svg_path"), width_inch=q.get("img_width_inch", 3.2))

        render_options(doc, q["options"], q.get("options_type", "text"))

    # 分頁後放置簡答與詳解
    doc.add_page_break()

    p_sol_title = doc.add_paragraph()
    p_sol_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sol_title.paragraph_format.space_after = Pt(4)
    r_sol_title = p_sol_title.add_run(f"{title} 參考解答與詳細解析")
    set_run_font(r_sol_title, font_name="標楷體", ascii_font="Times New Roman", size_pt=16, bold=True, color_rgb=COLOR_RED)

    # 簡答速查表
    table = doc.add_table(rows=2, cols=len(questions_data) + 1)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    headers = ["題號"] + [str(q["num"]) for q in questions_data]
    ans_row = ["答案"] + [f"({q['ans']})" if not q['ans'].startswith('(') else q['ans'] for q in questions_data]
    
    for col_idx, text in enumerate(headers):
        cell = table.cell(0, col_idx)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, font_name="標楷體", ascii_font="Times New Roman", size_pt=11.5, bold=True, color_rgb=COLOR_NAVY)
        
    for col_idx, text in enumerate(ans_row):
        cell = table.cell(1, col_idx)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, font_name="標楷體", ascii_font="Times New Roman", size_pt=12, bold=True, color_rgb=COLOR_RED)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(10)

    # 各題詳細解析 (全紅色)
    for q in questions_data:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        r_q = p_q.add_run(f"第 {q['num']} 題")
        set_run_font(r_q, font_name="標楷體", ascii_font="Times New Roman", size_pt=13, bold=True, color_rgb=COLOR_RED)
        
        # 答案與解析步驟
        for exp_line in q["explanation"]:
            add_runs_paragraph(doc, [exp_line], space_before=1, space_after=2, left_indent=0.2, force_color=COLOR_RED)

        if q.get("sol_image_path") and os.path.exists(q["sol_image_path"]):
            p_img = doc.add_paragraph()
            p_img.paragraph_format.left_indent = Inches(0.4)
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(4)
            add_exam_image(p_img, q["sol_image_path"], q.get("sol_svg_path"), width_inch=q.get("sol_img_width_inch", 3.2))

    ensure_document_font_consistency(doc, size_pt=13)
    doc.save(out_path)
    return out_path

def build_teacher_exam(questions_data, title, subtitle, out_path):
    """產出檔案 2：逐題詳解教師備課卷(B4_13pt版).docx"""
    doc = Document()
    configure_document_normal_style(doc, size_pt=13)
    upgrade_to_modern_word_mode(doc)
    configure_b4_section(doc.sections[0])

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run(f"{title}（教師備課逐題詳解卷）")
    set_run_font(r_title, font_name="標楷體", ascii_font="Times New Roman", size_pt=18, bold=True, color_rgb=COLOR_NAVY)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run(f"{subtitle}  ｜  全卷逐題對照詳解")
    set_run_font(r_sub, font_name="標楷體", ascii_font="Times New Roman", size_pt=12, color_rgb=COLOR_SUBTITLE)

    for idx, q in enumerate(questions_data):
        # 題幹 (黑字)
        add_runs_paragraph(doc, [q["stem"]], space_before=2, space_after=3, keep_with_next=True)

        # 附圖
        if q.get("image_path") and os.path.exists(q["image_path"]):
            p_img = doc.add_paragraph()
            p_img.paragraph_format.left_indent = Inches(0.4)
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.paragraph_format.keep_with_next = True
            add_exam_image(p_img, q["image_path"], q.get("svg_path"), width_inch=q.get("img_width_inch", 3.2))

        # 選項 (黑字)
        render_options(doc, q["options"], q.get("options_type", "text"))

        # 詳解 (全深紅字)
        for exp_line in q["explanation"]:
            add_runs_paragraph(doc, [exp_line], space_before=1, space_after=2, left_indent=0.2, force_color=COLOR_RED)

        # 詳解附圖
        if q.get("sol_image_path") and os.path.exists(q["sol_image_path"]):
            p_img = doc.add_paragraph()
            p_img.paragraph_format.left_indent = Inches(0.4)
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(4)
            add_exam_image(p_img, q["sol_image_path"], q.get("sol_svg_path"), width_inch=q.get("sol_img_width_inch", 3.2))

        # 題目間分隔線
        if idx < len(questions_data) - 1:
            p_div = doc.add_paragraph()
            p_div.paragraph_format.space_before = Pt(4)
            p_div.paragraph_format.space_after = Pt(8)
            r_div = p_div.add_run("―" * 58)
            set_run_font(r_div, font_name="標楷體", size_pt=10, color_rgb=COLOR_GRAY)

    ensure_document_font_consistency(doc, size_pt=13)
    doc.save(out_path)
    return out_path
